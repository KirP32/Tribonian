import base64
import io
import json
import zipfile
from collections.abc import Iterator
from pathlib import Path
from threading import Event

import pypdfium2 as pdfium
from bs4 import BeautifulSoup
from charset_normalizer import from_bytes
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation

from tribonian.domain import Unit, checkpoint

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
    ".log",
    ".rst",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
CHUNK_CHARS = 80000
MAX_FILE_BYTES = 100 * 1024 * 1024


def chunks(text: str, size: int = CHUNK_CHARS) -> Iterator[str]:
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind("\n", start + size // 2, end)
            if boundary != -1:
                end = boundary + 1
        yield text[start:end]
        start = end


def image_data(image: Image.Image) -> str:
    image = ImageOps.exif_transpose(image)
    image.thumbnail((1800, 1800))
    if image.mode in ("RGBA", "LA") or "transparency" in image.info:
        background = Image.new("RGB", image.size, "white")
        rgba = image.convert("RGBA")
        background.paste(rgba, mask=rgba.getchannel("A"))
        image = background
    else:
        image = image.convert("RGB")
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=90)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def text_units(text: str, label: str) -> Iterator[Unit]:
    for index, chunk in enumerate(chunks(text), 1):
        if chunk.strip():
            yield Unit(f"{label}, часть {index}", chunk)


def archive_images(path: Path, prefix: str, cancel: Event) -> Iterator[Unit]:
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            checkpoint(cancel)
            if name.startswith(prefix) and Path(name).suffix.lower() in IMAGE_EXTENSIONS:
                with Image.open(io.BytesIO(archive.read(name))) as image:
                    yield Unit(f"Встроенное изображение {name}", images=[image_data(image)])


def check_office_omissions(path: Path) -> None:
    problems = set()
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if any("/embeddings/" in name for name in names):
            problems.add("встроенные OLE-объекты")
        if any("/diagrams/" in name for name in names):
            problems.add("SmartArt")
        if any(name.startswith(("xl/charts/", "word/charts/")) for name in names):
            problems.add("визуальное представление диаграмм")
        if any("/media/" in name and Path(name).suffix.lower() not in IMAGE_EXTENSIONS for name in names):
            problems.add("векторные или неподдерживаемые медиавложения")
        if "word/document.xml" in names:
            document = archive.read("word/document.xml")
            for marker, description in [
                (b"<w:footnoteReference", "сноски"),
                (b"<w:endnoteReference", "концевые сноски"),
                (b"<w:txbxContent", "текстовые блоки в фигурах"),
            ]:
                if marker in document:
                    problems.add(description)
        for name in names:
            if name.startswith(("word/header", "word/footer")) and name.endswith(".xml"):
                if b"<w:tbl>" in archive.read(name):
                    problems.add("таблицы в колонтитулах")
    if problems:
        raise ValueError("Не извлечены: " + ", ".join(sorted(problems)) + ". Экспортируйте файл в PDF.")


def slide_shapes(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            yield from slide_shapes(shape.shapes)
        else:
            yield shape


def pdf_units(path: Path, cancel: Event) -> Iterator[Unit]:
    pending = ""
    first_page = 1
    with pdfium.PdfDocument(path) as document:
        for index in range(len(document)):
            checkpoint(cancel)
            page = document[index]
            try:
                text_page = page.get_textpage()
                try:
                    text = text_page.get_text_bounded()
                finally:
                    text_page.close()
                page_number = index + 1
                object_types = [obj.type for obj in page.get_objects(max_depth=3)]
                visual = (
                    len(text.strip()) < 80
                    or pdfium.raw.FPDF_PAGEOBJ_IMAGE in object_types
                    or object_types.count(pdfium.raw.FPDF_PAGEOBJ_PATH) > 10
                )
                if visual:
                    if pending:
                        yield from text_units(pending, f"PDF, стр. {first_page}–{index}")
                        pending = ""
                    width, height = page.get_size()
                    bitmap = page.render(scale=min(2, 1800 / max(width, height)))
                    try:
                        with bitmap.to_pil() as image:
                            data = image_data(image)
                    finally:
                        bitmap.close()
                    yield Unit(f"PDF, стр. {page_number}", text, [data])
                    first_page = page_number + 1
                else:
                    if pending and len(pending) + len(text) > CHUNK_CHARS:
                        yield from text_units(pending, f"PDF, стр. {first_page}–{index}")
                        pending = ""
                    if not pending:
                        first_page = page_number
                    pending += f"\n[Страница {page_number}]\n{text}\n"
            finally:
                page.close()
        if pending:
            yield from text_units(pending, f"PDF, стр. {first_page}–{len(document)}")


def extract_units(path: Path, cancel: Event) -> Iterator[Unit]:
    checkpoint(cancel)
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Файл больше 100 МБ. Разделите его на несколько документов.")
    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        raw = path.read_bytes()
        if not raw:
            return
        detected = from_bytes(raw).best()
        if detected is None:
            raise ValueError("Не удалось определить кодировку текста.")
        text = str(detected)
        if extension in {".html", ".htm"}:
            soup = BeautifulSoup(text, "html.parser")
            for element in soup(["script", "style"]):
                element.decompose()
            text = soup.get_text("\n", strip=True)
        if extension == ".json":
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        yield from text_units(text, "Текст")
    elif extension == ".pdf":
        yield from pdf_units(path, cancel)
    elif extension in IMAGE_EXTENSIONS:
        with Image.open(path) as image:
            for frame in range(getattr(image, "n_frames", 1)):
                checkpoint(cancel)
                image.seek(frame)
                yield Unit(f"Изображение, кадр {frame + 1}", images=[image_data(image.copy())])
    elif extension == ".docx":
        document = Document(path)
        lines = []
        for block in document.iter_inner_content():
            checkpoint(cancel)
            if hasattr(block, "text"):
                lines.append(block.text)
            else:
                for row in block.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
        for section in document.sections:
            for area in (section.header, section.footer):
                lines.extend(paragraph.text for paragraph in area.paragraphs)
        yield from text_units("\n".join(lines), "Word")
        yield from archive_images(path, "word/media/", cancel)
        check_office_omissions(path)
    elif extension == ".pptx":
        presentation = Presentation(path)
        for index, slide in enumerate(presentation.slides, 1):
            checkpoint(cancel)
            lines = []
            for shape in slide_shapes(slide.shapes):
                if shape.has_text_frame:
                    lines.append(shape.text)
                if shape.has_table:
                    lines.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
                if shape.has_chart:
                    lines.append("Диаграмма: " + str(shape.chart.chart_type))
                    for plot in shape.chart.plots:
                        categories = getattr(plot, "categories", [])
                        lines.append(
                            "Категории: " + ", ".join(str(category.label) for category in categories)
                        )
                    for series in shape.chart.series:
                        lines.append(f"{series.name}: {list(series.values)}")
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes:
                    lines.append(notes.text)
            yield from text_units("\n".join(lines), f"Слайд {index}")
        yield from archive_images(path, "ppt/media/", cancel)
        check_office_omissions(path)
    elif extension in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            for sheet in workbook:
                checkpoint(cancel)
                pending = ""
                for row in sheet.iter_rows():
                    checkpoint(cancel)
                    populated = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                    if populated:
                        pending += " | ".join(populated) + "\n"
                    if len(pending) >= CHUNK_CHARS:
                        yield from text_units(pending, f"Лист {sheet.title}")
                        pending = ""
                yield from text_units(pending, f"Лист {sheet.title}")
        finally:
            workbook.close()
        yield from archive_images(path, "xl/media/", cancel)
        check_office_omissions(path)
    else:
        raise ValueError(f"Формат {extension or 'без расширения'} пока не поддерживается.")
