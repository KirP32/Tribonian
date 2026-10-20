from pathlib import Path
from threading import Event

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdfium2 import PdfiumError

from tribonian.domain import Cancelled
from tribonian.extract import chunks, extract_units


def contents(path: Path) -> str:
    return "\n".join(unit.text for unit in extract_units(path, Event()))


@pytest.mark.parametrize("text", ["", "word " * 50000, "строка\n" * 30000, "x" * 160001])
def test_chunking_preserves_every_character(text):
    result = list(chunks(text))
    assert "".join(result) == text
    assert all(len(chunk) <= 80000 for chunk in result)


def test_cyrillic_encoding(tmp_path):
    path = tmp_path / "text.txt"
    path.write_bytes("Документ о гражданских правах и обязанностях.".encode("cp1251"))
    assert "гражданских правах" in contents(path)


def test_html_does_not_include_scripts(tmp_path):
    path = tmp_path / "page.html"
    path.write_text("<h1>Отчёт</h1><script>do_bad_things()</script><p>Доход: 42</p>", encoding="utf-8")
    text = contents(path)
    assert "Отчёт" in text and "Доход: 42" in text
    assert "do_bad_things" not in text


def test_docx_paragraphs_tables_and_images(tmp_path):
    image = tmp_path / "chart.png"
    Image.new("RGB", (120, 80), "green").save(image)
    document = Document()
    document.add_paragraph("План проекта")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Бюджет"
    table.cell(0, 1).text = "120000"
    document.add_picture(str(image))
    path = tmp_path / "plan.docx"
    document.save(path)
    units = list(extract_units(path, Event()))
    text = "\n".join(unit.text for unit in units)
    assert "План проекта" in text and "120000" in text
    assert any(unit.images for unit in units)


def test_xlsx_reads_all_sheets_and_formulas(tmp_path):
    workbook = Workbook()
    workbook.active["A1"] = "Выручка"
    workbook.active["B1"] = 300
    other = workbook.create_sheet("План")
    other["A1"] = "=Sheet!B1*2"
    path = tmp_path / "budget.xlsx"
    workbook.save(path)
    text = contents(path)
    assert "Выручка" in text and "300" in text and "=Sheet!B1*2" in text


def test_pptx_text_and_table(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Итоги квартала"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Рост"
    table.cell(0, 1).text = "25%"
    path = tmp_path / "slides.pptx"
    presentation.save(path)
    text = contents(path)
    assert "Итоги квартала" in text and "25%" in text


def test_image_and_scanned_pdf_use_vision(tmp_path):
    image = Image.new("RGB", (100, 120), "white")
    for suffix in [".png", ".pdf"]:
        path = tmp_path / ("scan" + suffix)
        image.save(path)
        units = list(extract_units(path, Event()))
        assert len(units) == 1
        assert units[0].images[0].startswith("data:image/jpeg;base64,")


def test_multipage_tiff_reads_all_frames(tmp_path):
    image = Image.new("RGB", (80, 80), "blue")
    path = tmp_path / "scan.tiff"
    image.save(path, save_all=True, append_images=[Image.new("RGB", (80, 80), "red")])
    assert len(list(extract_units(path, Event()))) == 2


def test_corrupt_pdf_is_not_treated_as_text(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_text("This is not a PDF.", encoding="utf-8")
    with pytest.raises(PdfiumError):
        list(extract_units(path, Event()))


def test_unknown_format_is_explicit(tmp_path):
    path = tmp_path / "notes.xyz"
    path.write_text("Unknown format", encoding="utf-8")
    with pytest.raises(ValueError, match="не поддерживается"):
        list(extract_units(path, Event()))


def test_cancel_prevents_extraction(tmp_path):
    stop = Event()
    stop.set()
    with pytest.raises(Cancelled):
        list(extract_units(tmp_path / "missing.pdf", stop))


def test_unrendered_chart_is_reported_after_extracting_cells(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Revenue", 42])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=1))
    sheet.add_chart(chart, "D1")
    path = tmp_path / "chart.xlsx"
    workbook.save(path)
    units = extract_units(path, Event())
    assert "42" in next(units).text
    with pytest.raises(ValueError, match="диаграмм"):
        list(units)
